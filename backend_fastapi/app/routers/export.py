from __future__ import annotations

from calendar import monthrange
from collections import defaultdict
from datetime import date
from io import BytesIO
import re
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_session
from app.security import require_roles

router = APIRouter()

HEADERS = ["Číslo výkazu", "Uživatel", "Zaměstnanec", "Druh výkazu", "Datum", "Od", "Do", "Pauza", "Hodiny", "Počet ha", "Kód stroje", "Traktor práce", "Pole", "Typ práce", "Popis jiné práce", "Středisko", "Poznámka"]
EXPORT_ROLES = ("admin", "reditel", "approved_viewer")


def cell(value) -> str:
    text_value = str(value or "")
    if isinstance(value, str) and text_value.startswith(("=", "+", "-", "@", "\t", "\r")):
        text_value = f"'{text_value}"
    return f'"{text_value.replace(chr(34), chr(34) + chr(34))}"'


def export_period(year: int | None, month: int | None) -> tuple[date | None, date | None]:
    if month is not None and year is None:
        raise HTTPException(status_code=422, detail="Pro výběr měsíce je potřeba zvolit rok.")
    if year is None:
        return None, None
    start = date(year, month or 1, 1)
    if month is None:
        return start, date(year + 1, 1, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return start, end


def export_filename(year: int | None, month: int | None, center: str) -> str:
    period = str(year) if year else "vse"
    if year and month:
        period = f"{year}_{month:02d}"
    normalized_center = unicodedata.normalize("NFKD", center).encode("ascii", "ignore").decode("ascii")
    center_slug = re.sub(r"[^a-z0-9]+", "_", normalized_center.casefold()).strip("_") or "vse"
    return f"adw_schvalene_{period}_{center_slug}.csv"


def safe_filename_part(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9]+", ".", normalized).strip(".") or "zamestnanec"


def workbook_filename(employee_name: str, year: int, month: int) -> str:
    return f"{safe_filename_part(employee_name)}.{month:02d}-{year}.xlsx"


def excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value if value is not None else ""


def workbook_fields(row: dict) -> str:
    entries = row.get("field_entries") or []
    if isinstance(entries, list):
        labels = []
        for entry in entries:
            if not isinstance(entry, dict) or not entry.get("field_name"):
                continue
            area = entry.get("amount_ha")
            labels.append(f"{entry['field_name']} ({area} ha)" if area not in (None, "") else str(entry["field_name"]))
        if labels:
            return "; ".join(labels)
    return str(row.get("field_name") or "")


def build_employee_workbook(employee_name: str, report_rows: list[dict], year: int, month: int) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows_by_day: dict[int, list[dict]] = defaultdict(list)
    for row in report_rows:
        rows_by_day[row["date"].day].append(row)

    sheet_headers = ["Číslo výkazu", "Od", "Do", "Hodiny", "Typ práce", "Popis jiné práce", "Středisko", "Pozemek", "Počet ha", "Stroj", "Poznámka"]
    for day in range(1, monthrange(year, month)[1] + 1):
        sheet = workbook.create_sheet(f"{day:02d}.{month:02d}.{year}")
        sheet.append([employee_name, f"{day:02d}.{month:02d}.{year}"])
        sheet.append(sheet_headers)
        for cell in sheet[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="005B4F")
        for row in rows_by_day.get(day, []):
            sheet.append([
                excel_value(row["report_number"]),
                str(row["time_start"] or "")[:5],
                str(row["time_end"] or "")[:5],
                row["hours_worked"],
                excel_value(row["work_type"]),
                excel_value(row["custom_work_description"]),
                excel_value(row["service_center"]),
                excel_value(workbook_fields(row)),
                row["amount_ha"],
                excel_value(row["tractor_name"]),
                excel_value(str(row["notes"] or "").replace("\n", " ")),
            ])
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:K{max(2, sheet.max_row)}"
        widths = [20, 8, 8, 10, 28, 38, 24, 24, 12, 28, 45]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get("/csv")
async def export_csv(
    center: str = "all",
    year: int | None = Query(default=None, ge=2000, le=2100),
    month: int | None = Query(default=None, ge=1, le=12),
    session: AsyncSession = Depends(get_session),
    _user=Depends(require_roles(*EXPORT_ROLES)),
):
    date_from, date_to = export_period(year, month)
    where = "WHERE r.status = 'approved' AND r.archived_at IS NULL"
    params = {}
    if center != "all":
        where += " AND r.service_center = :center"
        params["center"] = center
    if date_from is not None and date_to is not None:
        where += " AND r.date >= :date_from AND r.date < :date_to"
        params.update({"date_from": date_from, "date_to": date_to})
    result = await session.execute(
        text(
            f"""
            SELECT r.report_number, u.username, r.employee_name, r.report_kind, r.date, r.time_start, r.time_end, r.break_hours, r.hours_worked, r.amount_ha,
              t.tractor_code, t.tractor_name, f.field_name, w.name AS work_type, r.custom_work_description, r.service_center, r.notes
            FROM reports r
            LEFT JOIN users u ON r.user_id = u.id
            LEFT JOIN tractors t ON r.tractor_id = t.id
            LEFT JOIN fields f ON r.field_id = f.id
            LEFT JOIN work_types w ON r.work_type_id = w.id
            {where}
            ORDER BY r.date DESC
            """
        ),
        params,
    )
    rows = [HEADERS]
    for row in result.mappings().all():
        rows.append([
            row["report_number"], row["username"], row["employee_name"], row["report_kind"], str(row["date"])[:10], row["time_start"], row["time_end"], row["break_hours"],
            row["hours_worked"], row["amount_ha"], row["tractor_code"], row["tractor_name"], row["field_name"], row["work_type"], row["custom_work_description"],
            row["service_center"], str(row["notes"] or "").replace("\n", " "),
        ])
    csv = "\ufeff" + "\r\n".join(",".join(cell(item) for item in row) for row in rows)
    filename = export_filename(year, month, center)
    return Response(csv, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/xlsx")
async def export_xlsx(
    center: str = "all",
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    session: AsyncSession = Depends(get_session),
    _user=Depends(require_roles(*EXPORT_ROLES)),
):
    date_from, date_to = export_period(year, month)
    where = "WHERE r.status = 'approved' AND r.archived_at IS NULL AND r.date >= :date_from AND r.date < :date_to"
    params = {"date_from": date_from, "date_to": date_to}
    if center != "all":
        where += " AND r.service_center = :center"
        params["center"] = center
    result = await session.execute(
        text(
            f"""
            SELECT r.report_number, r.user_id, u.username, r.employee_name, r.report_kind, r.date,
              r.time_start, r.time_end, r.hours_worked, r.amount_ha, r.field_entries, t.tractor_name,
              f.field_name, w.name AS work_type, r.custom_work_description, r.service_center, r.notes
            FROM reports r
            LEFT JOIN users u ON r.user_id = u.id
            LEFT JOIN tractors t ON r.tractor_id = t.id
            LEFT JOIN fields f ON r.field_id = f.id
            LEFT JOIN work_types w ON r.work_type_id = w.id
            {where}
            ORDER BY r.employee_name, r.date, r.time_start, r.created_at
            """
        ),
        params,
    )
    grouped: dict[tuple[int, str], list[dict]] = defaultdict(list)
    for row in result.mappings().all():
        item = dict(row)
        grouped[(int(item["user_id"]), str(item["employee_name"] or item["username"] or "Zaměstnanec"))].append(item)
    if not grouped:
        raise HTTPException(status_code=404, detail="Pro zvolené období nejsou žádné schválené výkazy.")

    archive = BytesIO()
    with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
        for (_, employee_name), employee_rows in grouped.items():
            bundle.writestr(
                workbook_filename(employee_name, year, month),
                build_employee_workbook(employee_name, employee_rows, year, month),
            )
    center_suffix = safe_filename_part(center) if center != "all" else "vse"
    filename = f"adw_vykazy_{year}_{month:02d}_{center_suffix}.zip"
    return Response(
        archive.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
